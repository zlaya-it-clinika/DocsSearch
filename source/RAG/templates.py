"""
Этап 4. Заполнение модульного документа через диалог.

Логика (без ручной разметки шаблонов — новый модуль подхватывается автоматически):
  1. Вопросы строятся из «Примечаний по заполнению» модуля (data/templates/*.json):
     название поля + инструкция. Поля «Повторите ...» заполняются сами из ранее
     введённых, необязательные можно пропустить, строки таблиц состава — отдельные вопросы.
  2. Сборка .docx по правилам п. 3.3 Методических рекомендаций:
     удаляются гриф утверждения, пометка «МОДУЛЬНЫЙ ДОКУМЕНТ», примечания и номера [N];
     служебные зоны «НЕ УДАЛЯТЬ» сохраняются.
     Простые поля (одна линия — один номер) подставляет код. Сложные абзацы
     (варианты «проекту / проектам», несколько дат, необязательные части в [скобках])
     дописывает модель по собранным ответам, после чего код проверяет, что не осталось
     линий и номеров примечаний.
  3. Маршрут визирования — по Матрице (data/matrix.json), с уточняющим вопросом,
     если строка Матрицы содержит условие.
"""
import copy
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from config import DATA_DIR, SOURCE_DIR, LLM_MODEL, OLLAMA_HOST

OUT_DIR = DATA_DIR / "out"

MARKER_RE = re.compile(r"\s?\[(\d+)\]")
DATE_BLANK = r"«_+»\s*_{3,}(?:\s*20_+\s*г\.?)?"      # «___» ________ 20__ г. (год может быть один на две даты)
BLANK_RE = re.compile(rf"{DATE_BLANK}|«_+»|_{{3,}}|И\. О\. Фамилия")
DATE_RE = re.compile(DATE_BLANK)


def is_date_blank(s: str) -> bool:
    return bool(DATE_RE.fullmatch(s))
VARIANT_RE = re.compile(r"\[?([А-Яа-яЁё]+)\s*/\s*([А-Яа-яЁё]+)\]?")    # «[проекту / проектам]», «разработанного / измененного»
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]
KEEP_MARK = "НЕ УДАЛЯТЬ"


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s)


# ---------------------------------------------------------------- справочник модулей
def list_templates() -> list[dict]:
    out = []
    for f in sorted((DATA_DIR / "templates").glob("*.json")):
        t = json.loads(f.read_text(encoding="utf-8"))
        t["id"] = nfc(f.stem)
        out.append(t)
    return out


def template_title(t: dict) -> str:
    title = re.sub(r"\s*_{3,}.*$", "", t["title_pattern"]).strip()
    title = MARKER_RE.sub("", title)
    return f"{t['kind'].capitalize()} «{title} …»" if "_" in t["title_pattern"] else f"{t['kind'].capitalize()} «{title}»"


def find_docx(t: dict) -> Path:
    for f in SOURCE_DIR.glob("*.docx"):
        if nfc(f.name) == nfc(t["source"]):
            return f
    raise FileNotFoundError(f"Не найден файл модуля {t['source']} в {SOURCE_DIR}")


# ---------------------------------------------------------------- даты
def format_date(value: str) -> str:
    """01.11.2026 / 1.11.26 / 2026-11-01 -> 01 ноября 2026 г. (п. 8.7 Инструкции: форматы
    «14.10.2026» и «14 октября 2026 г.»); иначе как есть."""
    v = value.strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            d = datetime.strptime(v, fmt)
            return f"{d.day:02d} {MONTHS[d.month - 1]} {d.year} г."
        except ValueError:
            pass
    return v


# ---------------------------------------------------------------- обход документа
def iter_blocks(doc):
    """Абзацы и таблицы в порядке следования в документе."""
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            yield Paragraph(el, doc)
        elif tag == "tbl":
            yield Table(el, doc)


def set_paragraph_text(p: Paragraph, text: str):
    """Записывает текст в абзац, сохраняя оформление первого фрагмента; цвет — обычный."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    lines = text.split("\n")
    first = runs[0]
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    first.text = lines[0]
    first.font.color.rgb = None
    for line in lines[1:]:
        first.add_break()
        first.add_text(line)


def is_signature_table(tb: Table) -> bool:
    txt = " ".join(c.text for r in tb.rows for c in r.cells)
    return "Фамилия" in txt and ("Председател" in txt or "ВИЗУАЛИЗАЦИЯ" in txt)


def is_approval_table(tb: Table) -> bool:
    return "УТВЕРЖДАЮ" in " ".join(c.text for r in tb.rows for c in r.cells)


def composition_tables(doc) -> list[Table]:
    """Таблицы состава / назначаемых лиц: есть линии, не подпись и не гриф."""
    out = []
    for b in iter_blocks(doc):
        if isinstance(b, Table) and not is_signature_table(b) and not is_approval_table(b):
            if any(BLANK_RE.search(c.text) for r in b.rows for c in r.cells):
                out.append(b)
    return out


# ---------------------------------------------------------------- вопросы
class FillSession:
    """Диалог заполнения одного модуля. Используется из консоли (fill.py) и из API."""

    def __init__(self, template: dict, llm=None):
        self.t = template
        self.docx_path = find_docx(template)
        self.doc = Document(self.docx_path)
        self.llm = llm
        self.answers: dict[int, str | None] = {}     # номер поля -> значение (None — пропущено)
        self.table_answers: dict[str, str] = {}      # "т0:р1" -> значение
        self.questions = self._build_questions()
        self.pos = 0
        self.log: list[str] = []

    # --- построение вопросов
    def _body_text(self) -> str:
        parts = []
        for b in iter_blocks(self.doc):
            if isinstance(b, Paragraph):
                if b.text.startswith("Примечания по заполнению"):
                    break
                parts.append(b.text)
            else:
                parts += [c.text for r in b.rows for c in r.cells]
        return "\n".join(parts)

    def _build_questions(self) -> list[dict]:
        body = self._body_text()
        fields = self.t["fields"]
        by_id = {f["id"]: f for f in fields}
        in_body = {int(n) for n in MARKER_RE.findall(body)}

        tables = composition_tables(self.doc)
        table_field = None
        for f in fields:   # поле таблицы: его номер стоит в таблице или нигде в тексте
            in_tables = any(f"[{f['id']}]" in c.text for tb in tables for r in tb.rows for c in r.cells)
            if tables and (in_tables or (f["id"] not in in_body and f["name"] != "Подписант")):
                table_field = f["id"]
                break

        qs = []
        for f in fields:
            instr = f["instruction"]
            optional = (not f["required"]) or bool(re.search(r"только тогда|при необходимости отменить", instr))
            source = self._repeat_source(f, by_id)
            if f["id"] == table_field:
                qs += self._table_questions(tables, f)
                continue
            qs.append({"kind": "field", "field": f["id"], "name": f["name"], "hint": instr,
                       "optional": optional, "copy_from": source,
                       "is_date": bool(re.search(r"\bдат", f["name"], re.I))})
        return qs

    @staticmethod
    def _repeat_source(f: dict, by_id: dict) -> int | None:
        """«Повторите сокращенное наименование…», «Укажите те же…» -> из какого поля взять."""
        instr = f["instruction"].lower()
        if not re.search(r"повторите|те же|точно так же", instr):
            return None
        for other in by_id.values():
            if other["id"] < f["id"] and other["name"].lower() in instr:
                return other["id"]
        if "сокращенное наименование" in instr:
            for other in by_id.values():
                if "сокращенное" in other["name"].lower() and other["id"] < f["id"]:
                    return other["id"]
        return min(by_id)       # по умолчанию — первое поле (заголовок / предмет)

    def _table_questions(self, tables: list[Table], f: dict) -> list[dict]:
        qs = []
        for ti, tb in enumerate(tables):
            header = [c.text.strip() for c in tb.rows[0].cells]
            labelled = any(r.cells[0].text.strip() and not BLANK_RE.search(r.cells[0].text) for r in tb.rows[1:])
            if labelled:      # строки «Руководитель», «Секретарь», «Члены»
                for ri, row in enumerate(tb.rows[1:], 1):
                    label = row.cells[0].text.strip()
                    many = bool(re.search(r"Член|Аудитор", label))
                    qs.append({"kind": "table_row", "key": f"т{ti}:р{ri}", "table": ti, "row": ri,
                               "name": f"{f['name']}: {label}",
                               "hint": f["instruction"] + (" Несколько человек — через «;»." if many else ""),
                               "optional": False, "copy_from": None, "is_date": False})
            else:             # список: «Подразделение | Должность, ФИО» — строки добавляются по ответу
                qs.append({"kind": "table_list", "key": f"т{ti}", "table": ti,
                           "name": f["name"],
                           "hint": f["instruction"] + f" Каждую строку через «;», столбцы через « — »: "
                                                      f"{' — '.join(header)}.",
                           "optional": False, "copy_from": None, "is_date": False})
        return qs

    # --- диалог
    def next_question(self) -> dict | None:
        """Следующий вопрос; поля-повторы заполняются автоматически."""
        while self.pos < len(self.questions):
            q = self.questions[self.pos]
            if q["kind"] == "field" and q["copy_from"] is not None and self.answers.get(q["copy_from"]):
                self.answers[q["field"]] = self.answers[q["copy_from"]]
                self.log.append(f"Поле «{q['name']}» заполнено автоматически: «{self.answers[q['field']]}».")
                self.pos += 1
                continue
            return q
        return None

    def answer(self, text: str):
        q = self.questions[self.pos]
        text = text.strip()
        if not text and not q["optional"]:
            raise ValueError("Это поле обязательное.")
        value = text or None
        if q["kind"] == "field":
            self.answers[q["field"]] = format_date(value) if (value and q["is_date"]) else value
        else:
            self.table_answers[q["key"]] = value or ""
        self.pos += 1

    def progress(self) -> str:
        return f"{min(self.pos + 1, len(self.questions))}/{len(self.questions)}"

    # --- сборка документа
    def _data_lines(self) -> str:
        names = {f["id"]: f["name"] for f in self.t["fields"]}
        lines = []
        for n, v in sorted(self.answers.items()):
            lines.append(f"[{n}] {names.get(n, '')}: {v if v else 'НЕ ЗАПОЛНЯТЬ (необязательная часть — удалить)'}")
        return "\n".join(lines)

    def _simple_fill(self, text: str) -> str | None:
        """Подстановка без модели: каждый номер [N] — ровно одна соседняя линия, нет вариантов."""
        if "/" in BLANK_RE.sub("", text) or re.search(r"\[(?!\d+\])", text):
            return None
        markers = list(re.finditer(r"\[(\d+)\]", text))
        blanks = list(BLANK_RE.finditer(text))
        if not markers or len(markers) != len(blanks):
            return None
        out, last = [], 0
        for m in markers:
            n = int(m.group(1))
            value = self.answers.get(n)
            if value is None:
                return None
            before = [b for b in blanks if b.end() <= m.start() and b.start() >= last]
            after = [b for b in blanks if b.start() >= m.end()]
            gap_before = text[before[-1].end():m.start()] if before else None
            if before and len(gap_before) <= 3:
                b = before[-1]
            elif after and len(text[m.end():after[0].start()]) <= 3:
                b = after[0]
            else:
                return None
            v = format_date(value) if is_date_blank(b.group(0)) else value
            if b.group(0).startswith("«") and not is_date_blank(b.group(0)):
                v = "«" + v.strip("«»\"") + "»"           # линия в кавычках — кавычки сохраняем
            if v.endswith(".") and (text[b.end():].lstrip().startswith(".")
                                    or text[m.end():].lstrip().startswith(".")):
                v = v[:-1]                                # «Иванов И. И.» + «.» шаблона
            if b.start() < m.start():
                out.append(text[last:b.start()] + v + text[b.end():m.start()])
                last = m.end()
            else:
                out.append(text[last:m.start()] + text[m.end():b.start()] + v)
                last = b.end()
        out.append(text[last:])
        res = "".join(out)
        return re.sub(r"\s+([,.;])", r"\1", res).replace("  ", " ")

    def _slot_template(self, text: str) -> tuple[str, list[dict]]:
        """Превращает абзац в шаблон с пропусками {S1}, {S2}…: линии и варианты «а / б».

        Модель потом возвращает только значения пропусков (JSON), а текст абзаца собирает код —
        так модель физически не может переписать или дописать постоянные формулировки.
        """
        # заполненная необязательная часть: «[Признать утратившим силу ____.]» -> без скобок
        text = re.sub(r"\[([^\[\]/]*?_{3,}[^\[\]]*?)\]", r"\1", text)
        text = re.sub(r"\[(и об [^\[\]]*?)\]", r"\1", text)
        tokens = [(m.start(), m.end(), "blank", m) for m in BLANK_RE.finditer(text)]
        tokens += [(m.start(), m.end(), "variant", m) for m in VARIANT_RE.finditer(text)]
        tokens.sort(key=lambda x: x[0])
        markers = [(m.start(), int(m.group(1))) for m in re.finditer(r"\[(\d+)\]", text)]
        names = {f["id"]: f["name"] for f in self.t["fields"]}

        out, slots, last = [], [], 0
        for start, end, kind, m in tokens:
            if start < last:
                continue
            sid = f"S{len(slots) + 1}"
            slot = {"id": sid, "kind": kind, "raw": m.group(0)}
            if kind == "variant":
                slot["options"] = [m.group(1), m.group(2)]
            else:
                after = [n for pos, n in markers if pos >= end]
                before = [n for pos, n in markers if pos < start]
                n = after[0] if after else (before[-1] if before else None)
                slot["field"] = n
                slot["field_name"] = names.get(n, "")
                slot["is_date"] = is_date_blank(m.group(0))
                slot["quoted"] = m.group(0).startswith("«") and not slot["is_date"]
            out.append(text[last:start] + "{" + sid + "}")
            slots.append(slot)
            last = end
        out.append(text[last:])
        tpl = re.sub(r"\s?\[\d+\]", "", "".join(out))
        return tpl, slots

    def _ask_slots(self, tpl: str, slots: list[dict]) -> dict | None:
        lines = []
        for sl in slots:
            if sl["kind"] == "variant":
                lines.append(f'{sl["id"]} — выбрать один вариант: «{sl["options"][0]}» или «{sl["options"][1]}»')
            elif sl["field"]:
                kind = "дата, " if sl["is_date"] else ""
                lines.append(f'{sl["id"]} — {kind}данные поля [{sl["field"]}] «{sl["field_name"]}»')
            else:
                lines.append(f'{sl["id"]} — по смыслу, из данных полей')
        prompt = (
            "Заполни пропуски в абзаце официального документа.\n"
            f"АБЗАЦ:\n{tpl}\n\nПРОПУСКИ:\n" + "\n".join(lines) +
            f"\n\nДАННЫЕ ПОЛЕЙ:\n{self._data_lines()}\n\n"
            "ПРАВИЛА: для каждого пропуска возьми данные указанного поля. Можно изменить падеж и число "
            "слов, чтобы фраза в абзаце была грамматически правильной, но нельзя менять смысл и добавлять "
            "сведения, которых нет в данных. Если в одном поле несколько значений (например, две даты), "
            "раздели их по соответствующим пропускам. Для вариантов выбери тот, что подходит к данным. "
            "Даты пиши в формате ДД.ММ.ГГГГ (код сам переведёт их в формат «01 ноября 2026 г.»). "
            'Ответь только JSON без пояснений: {"S1": "...", "S2": "..."}'
        )
        for _ in range(2):
            r = self.llm.chat(model=LLM_MODEL, messages=[{"role": "user", "content": prompt}],
                              options={"temperature": 0.0, "num_ctx": 4096, "num_predict": 300})
            m = re.search(r"\{.*\}", r["message"]["content"], re.S)
            try:
                data = json.loads(m.group(0)) if m else None
            except json.JSONDecodeError:
                data = None
            if isinstance(data, dict):
                return {k: str(v) for k, v in data.items()}
        return None

    def _assemble(self, tpl: str, slots: list[dict], values: dict) -> str:
        res = tpl
        for sl in slots:
            v = (values.get(sl["id"]) or "").strip()
            if sl["kind"] == "variant":
                v = v.strip("«»") if v.strip("«»") in sl["options"] else sl["options"][0]
            else:
                if not v and sl["field"]:
                    v = self.answers.get(sl["field"]) or ""
                if sl["is_date"]:
                    v = format_date(v)
                if sl["quoted"]:
                    v = "«" + v.strip("«»\"") + "»"
            ph = "{" + sl["id"] + "}"
            tail = res[res.index(ph) + len(ph):] if ph in res else ""
            if v.endswith(".") and tail.lstrip().startswith("."):
                v = v[:-1]
            if tail.startswith("»") and not v.startswith("«") and "«" not in res[:res.index(ph)][-3:]:
                v = "«" + v                                   # «(далее - ____»)» в шаблоне без открывающей кавычки
            res = res.replace(ph, v, 1)
        res = re.sub(r"[ \t]+([,.;])", r"\1", res)
        return re.sub(r"[ \t]{2,}", " ", res)

    def _llm_fill(self, text: str) -> str:
        tpl, slots = self._slot_template(text)
        values = self._ask_slots(tpl, slots) if self.llm is not None else None
        if values is None:
            values = {}
            self.log.append("Модель не вернула значения — подставлены ответы как есть.")
        return self._assemble(tpl, slots, values)

    def _drop_skipped_optional(self, text: str) -> str:
        """Необязательная часть «[ ... ] [N]», если поле N пропущено, удаляется кодом, а не моделью."""
        def repl(m):
            n = int(m.group(1))
            return "" if (n in self.answers and self.answers[n] is None) else m.group(0)
        return re.sub(r"\n?\[[^\[\]]*\]\s*\[(\d+)\]", repl, text)

    def _fill_paragraph(self, p: Paragraph, problems: list[str]):
        original = p.text
        text = self._drop_skipped_optional(original)
        if re.fullmatch(r"\s*\d+\.\s*", text):            # необязательный пункт удалён целиком
            p._element.getparent().remove(p._element)
            return
        if KEEP_MARK in text or not (BLANK_RE.search(text) or MARKER_RE.search(text) or re.search(r"\[[^\]]*/", text)):
            if text != original:
                set_paragraph_text(p, text.strip())
            return
        result = self._simple_fill(text)
        if result is None:
            result = self._llm_fill(text)
            self.log.append(f"Абзац заполнен с помощью модели: «{result[:90]}…»")
        if BLANK_RE.search(result) or re.search(r"\[\d+\]", result):
            problems.append(f"Проверьте абзац: «{result[:100]}»")
        result = result.strip()
        if re.fullmatch(r"\d+\.\s*\.?", result) or not result:     # необязательный пункт удалён целиком
            p._element.getparent().remove(p._element)
            return
        set_paragraph_text(p, result)

    def _fill_tables(self, tables: list[Table]):
        for q in self.questions:
            if q["kind"] == "table_row":
                tb = tables[q["table"]]
                val = self.table_answers.get(q["key"], "")
                cell = tb.rows[q["row"]].cells[-1]
                set_paragraph_text(cell.paragraphs[0], "\n".join(s.strip() for s in val.split(";")))
            elif q["kind"] == "table_list":
                tb = tables[q["table"]]
                items = [s.strip() for s in self.table_answers.get(q["key"], "").split(";") if s.strip()]
                template_row = tb.rows[1]._tr
                for r in list(tb.rows[1:]):          # убираем пустые строки шаблона
                    tb._tbl.remove(r._tr)
                for item in items:
                    tr = copy.deepcopy(template_row)
                    tb._tbl.append(tr)
                    row = tb.rows[-1]
                    cols = [c.strip() for c in re.split(r"\s+[—–-]\s+", item, maxsplit=len(row.cells) - 1)]
                    for c, v in zip(row.cells, cols + [""] * len(row.cells)):
                        set_paragraph_text(c.paragraphs[0], v)

    def _fill_signature(self, tb: Table):
        names = {f["name"]: f["id"] for f in self.t["fields"]}
        signer = self.answers.get(names.get("Подписант"))
        for row in tb.rows:
            for c in row.cells:
                if "Фамилия" in c.text:
                    set_paragraph_text(c.paragraphs[0], signer or MARKER_RE.sub("", c.text))
                elif re.fullmatch(r"\s*_{3,}\s*", c.text):
                    # линия для собственноручной подписи: ЭД подписывается ЭП (п. 9.5.3, 8.19),
                    # линии удаляются (п. 3.3 Методических рекомендаций); зона визуализации ЭП остаётся
                    set_paragraph_text(c.paragraphs[0], "")

    def build(self) -> dict:
        """Собирает .docx. Возвращает путь, замечания и маршрут визирования."""
        doc = self.doc
        problems = []
        tables = composition_tables(doc)
        self._fill_tables(tables)

        removing = False
        for b in list(iter_blocks(doc)):
            el = b._element if isinstance(b, Paragraph) else b._tbl
            if isinstance(b, Paragraph) and b.text.startswith("Примечания по заполнению"):
                removing = True
            if removing:
                el.getparent().remove(el)
                continue
            if isinstance(b, Table):
                if is_approval_table(b):
                    el.getparent().remove(el)            # гриф утверждения
                elif is_signature_table(b):
                    self._fill_signature(b)
                continue
            if b.text.strip() == "МОДУЛЬНЫЙ ДОКУМЕНТ":
                el.getparent().remove(el)
                continue
            self._fill_paragraph(b, problems)

        self._renumber(doc)
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = OUT_DIR / f"{self.t['id'].replace('Модуль_', 'Проект_')}_{stamp}.docx"
        doc.save(out)
        return {"path": out, "problems": problems, "log": self.log}

    @staticmethod
    def _renumber(doc):
        """После удаления необязательного пункта — сквозная нумерация 1., 2., 3. …"""
        started, n = False, 0
        for b in iter_blocks(doc):
            if not isinstance(b, Paragraph):
                continue
            if b.text.strip() in ("ПРИКАЗЫВАЮ:", "ОБЯЗЫВАЮ:"):
                started = True
                continue
            m = re.match(r"^(\d+)\.\s", b.text)
            if started and m:
                n += 1
                if int(m.group(1)) != n and b.runs:
                    b.runs[0].text = re.sub(r"^\d+\.", f"{n}.", b.runs[0].text, count=1)


# ---------------------------------------------------------------- маршрут по Матрице
def matrix_route(t: dict) -> dict:
    """Строка Матрицы для модуля + условие, по которому нужно спросить пользователя."""
    matrix = json.loads((DATA_DIR / "matrix.json").read_text(encoding="utf-8"))
    rows = {r["id"]: r for r in matrix["rows"]}
    name = t["title_pattern"].lower() + " " + t["source"].lower()
    if "эксплуатаци" in name:
        rid = "ВН-3"
    elif "передач" in name and "дел" in name:
        rid = "ВН-2"
    elif re.search(r"рабоч\w* групп|проектн\w* комитет|комисси", name):
        rid = "ВН-6"
    elif t["kind"] == "ПРИКАЗ":
        rid = "ВН-4"
    else:
        rid = "ВН-7"
    row = rows[rid]
    always = [k for k, v in row["visas"].items() if v.strip().lower() == "да"]
    other = row["other"] or ""
    if other.strip().lower() in ("нет", "-", ""):
        other = ""
    condition = other if len(other.split()) > 8 else None
    if other and not condition:
        always += [s.strip() for s in other.split(",")]
    return {"row": rid, "doc_type": row["doc_type"], "always": always, "condition": condition}
